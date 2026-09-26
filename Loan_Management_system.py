import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os
# EXCEL FILE
EXCEL_FILE = "loan_records.xlsx"
def create_excel_file():
    if not os.path.exists(EXCEL_FILE):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Loans"

        sheet.append([
            "Loan ID",
            "Customer Name",
            "Mobile",
            "Loan Type",
            "Loan Amount",
            "Interest Rate",
            "Tenure",
            "EMI",
            "Status"
        ])

        workbook.save(EXCEL_FILE)

create_excel_file()

# MAIN WINDOW
window = tk.Tk()
window.title("Loan Management System")
window.geometry("950x650")
window.resizable(False, False)

# COMMON FUNCTIONS
def clear_form():
    name_entry.delete(0, tk.END)
    mobile_entry.delete(0, tk.END)
    loan_type_entry.delete(0, tk.END)
    amount_entry.delete(0, tk.END)
    interest_entry.delete(0, tk.END)
    tenure_entry.delete(0, tk.END)
    emi_entry.delete(0, tk.END)

    status_combo.set("Pending")

def calculate_emi():
    try:
        amount = float(amount_entry.get())
        annual_rate = float(interest_entry.get())
        months = int(tenure_entry.get())

        if amount <= 0 or annual_rate < 0 or months <= 0:
            raise ValueError

        monthly_rate = annual_rate / 12 / 100

        if monthly_rate == 0:
            emi = amount / months
        else:
            emi = (
                amount
                * monthly_rate
                * (1 + monthly_rate) ** months
                / ((1 + monthly_rate) ** months - 1)
            )

        emi_entry.delete(0, tk.END)
        emi_entry.insert(0, f"{emi:.2f}")

    except ValueError:
        messagebox.showerror(
            "Error",
            "Please enter valid Loan Amount, Interest Rate and Tenure."
        )


# LOGIN
def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "1234":
        login_frame.pack_forget()
        show_dashboard()

    else:
        messagebox.showerror(
            "Login Error",
            "Invalid Username or Password"
        )

# LOGOUT


def logout():
    dashboard_frame.pack_forget()
    add_frame.pack_forget()
    view_frame.pack_forget()
    search_frame.pack_forget()

    clear_form()

    username_entry.delete(0, tk.END)
    password_entry.delete(0, tk.END)

    login_frame.pack(fill="both", expand=True)



# DASHBOARD
def show_dashboard():

    add_frame.pack_forget()
    view_frame.pack_forget()
    search_frame.pack_forget()

    dashboard_frame.pack(fill="both", expand=True)



# ADD LOAN
def show_add_loan():

    dashboard_frame.pack_forget()
    view_frame.pack_forget()
    search_frame.pack_forget()

    add_frame.pack(fill="both", expand=True)


def save_loan():

    name = name_entry.get()
    mobile = mobile_entry.get()
    loan_type = loan_type_entry.get()
    amount = amount_entry.get()
    interest = interest_entry.get()
    tenure = tenure_entry.get()
    emi = emi_entry.get()
    status = status_combo.get()

    if (
        name == ""
        or mobile == ""
        or loan_type == ""
        or amount == ""
        or interest == ""
        or tenure == ""
    ):
        messagebox.showwarning(
            "Warning",
            "Please fill all required fields."
        )
        return

    try:
        float(amount)
        float(interest)
        int(tenure)

    except ValueError:
        messagebox.showerror(
            "Error",
            "Please enter valid numeric values."
        )
        return

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Loans"]

    loan_id = sheet.max_row

    sheet.append([
        loan_id,
        name,
        mobile,
        loan_type,
        amount,
        interest,
        tenure,
        emi,
        status
    ])

    workbook.save(EXCEL_FILE)

    messagebox.showinfo(
        "Success",
        "Loan record added successfully!"
    )

    clear_form()

# VIEW RECORDS
def show_view_records():

    dashboard_frame.pack_forget()
    add_frame.pack_forget()
    search_frame.pack_forget()

    view_frame.pack(fill="both", expand=True)

    load_records()


def load_records():

    for item in tree.get_children():
        tree.delete(item)

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Loans"]

    for row in sheet.iter_rows(min_row=2, values_only=True):

        tree.insert("", tk.END, values=row)

# SEARCH
def show_search():

    dashboard_frame.pack_forget()
    add_frame.pack_forget()
    view_frame.pack_forget()

    search_frame.pack(fill="both", expand=True)

    for item in search_tree.get_children():
        search_tree.delete(item)


def search_record():

    search_value = search_entry.get().strip().lower()

    for item in search_tree.get_children():
        search_tree.delete(item)

    if search_value == "":
        messagebox.showwarning(
            "Warning",
            "Enter Loan ID or Customer Name."
        )
        return

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Loans"]

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        loan_id = str(row[0]).lower()
        customer_name = str(row[1]).lower()

        if search_value in loan_id or search_value in customer_name:

            search_tree.insert(
                "",
                tk.END,
                values=row
            )

            found = True

    if not found:
        messagebox.showinfo(
            "Search",
            "No record found."
        )

# UPDATE
def update_record():

    selected = tree.selection()

    if not selected:
        messagebox.showwarning(
            "Warning",
            "Please select a record from the table."
        )
        return

    values = tree.item(selected[0], "values")

    loan_id = values[0]

    name_entry.delete(0, tk.END)
    name_entry.insert(0, values[1])

    mobile_entry.delete(0, tk.END)
    mobile_entry.insert(0, values[2])

    loan_type_entry.delete(0, tk.END)
    loan_type_entry.insert(0, values[3])

    amount_entry.delete(0, tk.END)
    amount_entry.insert(0, values[4])

    interest_entry.delete(0, tk.END)
    interest_entry.insert(0, values[5])

    tenure_entry.delete(0, tk.END)
    tenure_entry.insert(0, values[6])

    emi_entry.delete(0, tk.END)
    emi_entry.insert(0, values[7])

    status_combo.set(values[8])

    show_add_loan()

    update_button.pack(pady=5)

    update_button.config(
        command=lambda: save_update(loan_id)
    )


def save_update(loan_id):

    name = name_entry.get()
    mobile = mobile_entry.get()
    loan_type = loan_type_entry.get()
    amount = amount_entry.get()
    interest = interest_entry.get()
    tenure = tenure_entry.get()
    emi = emi_entry.get()
    status = status_combo.get()

    if name == "" or mobile == "" or loan_type == "" or amount == "":
        messagebox.showwarning(
            "Warning",
            "Please fill all required fields."
        )
        return

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Loans"]

    for row in sheet.iter_rows(min_row=2):

        if str(row[0].value) == str(loan_id):

            row[1].value = name
            row[2].value = mobile
            row[3].value = loan_type
            row[4].value = amount
            row[5].value = interest
            row[6].value = tenure
            row[7].value = emi
            row[8].value = status

            break

    workbook.save(EXCEL_FILE)

    messagebox.showinfo(
        "Success",
        "Loan record updated successfully!"
    )

    update_button.pack_forget()

    clear_form()

    show_view_records()


# DELETE


def delete_record():

    selected = tree.selection()

    if not selected:
        messagebox.showwarning(
            "Warning",
            "Please select a record to delete."
        )
        return

    values = tree.item(selected[0], "values")

    loan_id = values[0]

    confirm = messagebox.askyesno(
        "Confirm Delete",
        "Are you sure you want to delete this record?"
    )

    if not confirm:
        return

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Loans"]

    for row in range(2, sheet.max_row + 1):

        if str(sheet.cell(row, 1).value) == str(loan_id):

            sheet.delete_rows(row, 1)
            break

    workbook.save(EXCEL_FILE)

    messagebox.showinfo(
        "Success",
        "Loan record deleted successfully!"
    )

    load_records()



# LOGIN FRAME


login_frame = tk.Frame(
    window,
    bg="white"
)


title_label = tk.Label(
    login_frame,
    text="LOAN MANAGEMENT SYSTEM",
    font=("Arial", 24, "bold"),
    bg="white"
)

title_label.pack(pady=50)


tk.Label(
    login_frame,
    text="Username",
    font=("Arial", 12),
    bg="white"
).pack()


username_entry = tk.Entry(
    login_frame,
    width=30,
    font=("Arial", 12)
)

username_entry.pack(pady=8)


tk.Label(
    login_frame,
    text="Password",
    font=("Arial", 12),
    bg="white"
).pack()


password_entry = tk.Entry(
    login_frame,
    width=30,
    font=("Arial", 12),
    show="*"
)

password_entry.pack(pady=8)


tk.Button(
    login_frame,
    text="LOGIN",
    width=15,
    font=("Arial", 12, "bold"),
    command=login
).pack(pady=25)



# DASHBOARD FRAME


dashboard_frame = tk.Frame(
    window,
    bg="lightblue"
)


tk.Label(
    dashboard_frame,
    text="LOAN MANAGEMENT SYSTEM",
    font=("Arial", 25, "bold"),
    bg="lightblue"
).pack(pady=30)


tk.Label(
    dashboard_frame,
    text="DASHBOARD",
    font=("Arial", 20, "bold"),
    bg="lightblue"
).pack(pady=10)


tk.Button(
    dashboard_frame,
    text="ADD LOAN",
    width=25,
    command=show_add_loan
).pack(pady=6)


tk.Button(
    dashboard_frame,
    text="VIEW RECORDS",
    width=25,
    command=show_view_records
).pack(pady=6)


tk.Button(
    dashboard_frame,
    text="SEARCH",
    width=25,
    command=show_search
).pack(pady=6)


tk.Button(
    dashboard_frame,
    text="DELETE SELECTED",
    width=25,
    command=delete_record
).pack(pady=6)


tk.Button(
    dashboard_frame,
    text="LOGOUT",
    width=25,
    command=logout
).pack(pady=20)



# ADD FRAME


add_frame = tk.Frame(
    window,
    bg="white"
)


tk.Label(
    add_frame,
    text="ADD LOAN RECORD",
    font=("Arial", 22, "bold"),
    bg="white"
).grid(
    row=0,
    column=0,
    columnspan=2,
    pady=20
)


tk.Label(
    add_frame,
    text="Customer Name",
    bg="white"
).grid(row=1, column=0, padx=20, pady=8, sticky="w")

name_entry = tk.Entry(add_frame, width=30)
name_entry.grid(row=1, column=1, pady=8)


tk.Label(
    add_frame,
    text="Mobile Number",
    bg="white"
).grid(row=2, column=0, padx=20, pady=8, sticky="w")

mobile_entry = tk.Entry(add_frame, width=30)
mobile_entry.grid(row=2, column=1, pady=8)


tk.Label(
    add_frame,
    text="Loan Type",
    bg="white"
).grid(row=3, column=0, padx=20, pady=8, sticky="w")

loan_type_entry = tk.Entry(add_frame, width=30)
loan_type_entry.grid(row=3, column=1, pady=8)


tk.Label(
    add_frame,
    text="Loan Amount",
    bg="white"
).grid(row=4, column=0, padx=20, pady=8, sticky="w")

amount_entry = tk.Entry(add_frame, width=30)
amount_entry.grid(row=4, column=1, pady=8)


tk.Label(
    add_frame,
    text="Interest Rate (%)",
    bg="white"
).grid(row=5, column=0, padx=20, pady=8, sticky="w")

interest_entry = tk.Entry(add_frame, width=30)
interest_entry.grid(row=5, column=1, pady=8)


tk.Label(
    add_frame,
    text="Tenure (Months)",
    bg="white"
).grid(row=6, column=0, padx=20, pady=8, sticky="w")

tenure_entry = tk.Entry(add_frame, width=30)
tenure_entry.grid(row=6, column=1, pady=8)


tk.Label(
    add_frame,
    text="EMI",
    bg="white"
).grid(row=7, column=0, padx=20, pady=8, sticky="w")

emi_entry = tk.Entry(add_frame, width=30)
emi_entry.grid(row=7, column=1, pady=8)


tk.Label(
    add_frame,
    text="Status",
    bg="white"
).grid(row=8, column=0, padx=20, pady=8, sticky="w")


status_combo = ttk.Combobox(
    add_frame,
    values=["Pending", "Approved", "Rejected"],
    width=27,
    state="readonly"
)

status_combo.set("Pending")
status_combo.grid(row=8, column=1, pady=8)


tk.Button(
    add_frame,
    text="CALCULATE EMI",
    command=calculate_emi
).grid(row=9, column=0, pady=15)


tk.Button(
    add_frame,
    text="SAVE LOAN",
    command=save_loan
).grid(row=9, column=1, pady=15)


update_button = tk.Button(
    add_frame,
    text="UPDATE LOAN"
)


tk.Button(
    add_frame,
    text="CLEAR",
    command=clear_form
).grid(row=10, column=0, pady=10)


tk.Button(
    add_frame,
    text="BACK",
    command=show_dashboard
).grid(row=10, column=1, pady=10)



# VIEW FRAME


view_frame = tk.Frame(
    window,
    bg="white"
)


tk.Label(
    view_frame,
    text="LOAN RECORDS",
    font=("Arial", 22, "bold"),
    bg="white"
).pack(pady=15)


columns = (
    "Loan ID",
    "Customer Name",
    "Mobile",
    "Loan Type",
    "Loan Amount",
    "Interest Rate",
    "Tenure",
    "EMI",
    "Status"
)


tree = ttk.Treeview(
    view_frame,
    columns=columns,
    show="headings",
    height=18
)


for column in columns:

    tree.heading(
        column,
        text=column
    )

    tree.column(
        column,
        width=100
    )


tree.pack(
    padx=10,
    pady=10,
    fill="both"
)


button_frame = tk.Frame(
    view_frame,
    bg="white"
)

button_frame.pack(pady=10)


tk.Button(
    button_frame,
    text="UPDATE SELECTED",
    command=update_record
).grid(row=0, column=0, padx=5)


tk.Button(
    button_frame,
    text="DELETE SELECTED",
    command=delete_record
).grid(row=0, column=1, padx=5)


tk.Button(
    button_frame,
    text="REFRESH",
    command=load_records
).grid(row=0, column=2, padx=5)


tk.Button(
    button_frame,
    text="BACK",
    command=show_dashboard
).grid(row=0, column=3, padx=5)



# SEARCH FRAME


search_frame = tk.Frame(
    window,
    bg="white"
)


tk.Label(
    search_frame,
    text="SEARCH LOAN RECORD",
    font=("Arial", 22, "bold"),
    bg="white"
).pack(pady=20)


search_entry = tk.Entry(
    search_frame,
    width=40,
    font=("Arial", 12)
)

search_entry.pack(pady=10)


tk.Button(
    search_frame,
    text="SEARCH",
    width=15,
    command=search_record
).pack(pady=10)


search_tree = ttk.Treeview(
    search_frame,
    columns=columns,
    show="headings",
    height=15
)


for column in columns:

    search_tree.heading(
        column,
        text=column
    )

    search_tree.column(
        column,
        width=100
    )


search_tree.pack(
    padx=10,
    pady=15
)


tk.Button(
    search_frame,
    text="BACK",
    command=show_dashboard
).pack(pady=10)

# START APPLICATION

login_frame.pack(
    fill="both",
    expand=True
)


window.mainloop()